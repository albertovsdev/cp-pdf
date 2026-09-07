"""Objetivo 3 (fase 8d): el mismo patron en los otros cuatro exportadores.

La pregunta es una sola: hay alguna hoja que muestre una cifra DECLARADA
por el documento cuando el sistema tambien tiene la cifra LEIDA, sin
ensenar las dos? Y si la hay, difieren en algun fixture? Se mide; no se
arregla nada aqui.
"""

from __future__ import annotations

import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "tests"))

import openpyxl  # noqa: E402
from conftest import REAL_PDFS  # noqa: E402

from contapdf.export.excel import (  # noqa: E402
    exportar_auxiliar,
    exportar_balanza,
    exportar_estado_cuenta,
    exportar_mayor,
)
from contapdf.pipeline import (  # noqa: E402
    procesar_auxiliar,
    procesar_balanza,
    procesar_estado_cuenta,
    procesar_mayor,
)

_CERO = Decimal(0)
salida = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/tmp")


def _columnas(destino: Path) -> dict[str, list[str]]:
    libro = openpyxl.load_workbook(destino)
    return {h: [c.value for c in libro[h][1]] for h in libro.sheetnames}


def _bloque(titulo: str) -> None:
    print()
    print("=" * 74)
    print(titulo)
    print("=" * 74)


# --- MAYOR -------------------------------------------------------------
_bloque("mayor-gume: totales de la cuenta (declarados) contra la suma de meses")
r = procesar_mayor(REAL_PDFS["mayor-gume"])
destino = salida / "m6-mayor.xlsx"
exportar_mayor(r.mayor, r.cobertura, destino)
print("hojas y columnas:")
for hoja, cols in _columnas(destino).items():
    print(f"  {hoja:14} {cols}")
difieren = 0
for cuenta in r.mayor.cuentas:
    suyos = [m for m in r.mayor.meses if m.cuenta == cuenta.cuenta]
    if not suyos:
        continue
    suma_cargos = sum((m.cargos for m in suyos), _CERO)
    suma_abonos = sum((m.abonos for m in suyos), _CERO)
    if (cuenta.total_cargos is not None and cuenta.total_cargos != suma_cargos) \
            or (cuenta.total_abonos is not None
                and cuenta.total_abonos != suma_abonos):
        difieren += 1
print(f"cuentas                                   {len(r.mayor.cuentas)}")
print(f"cuentas donde el total declarado difiere  {difieren}")

# --- ESTADO DE CUENTA --------------------------------------------------
_bloque("estados de cuenta: el resumen (declarado) contra la suma de movimientos")
for nombre in ("edocta", "edocta-bbva", "edocta-julio-banorte", "edocta-bajio"):
    r = procesar_estado_cuenta(REAL_PDFS[nombre])
    destino = salida / f"m6-{nombre}.xlsx"
    exportar_estado_cuenta(r.estado, r.cobertura, destino)
    if nombre == "edocta":
        print("hojas y columnas:")
        for hoja, cols in _columnas(destino).items():
            print(f"  {hoja:14} {cols}")
    difieren = sin_desglose = 0
    for cuenta in r.estado.cuentas:
        suyos = r.estado.movimientos_de(cuenta.num_cuenta)
        if cuenta.depositos is None or cuenta.retiros is None:
            sin_desglose += 1
            continue
        dep = sum((m.deposito or _CERO for m in suyos), _CERO)
        ret = sum((m.retiro or _CERO for m in suyos), _CERO)
        if cuenta.depositos != dep or cuenta.retiros != ret:
            difieren += 1
    print(f"  {nombre:22} cuentas={len(r.estado.cuentas):2} "
          f"difieren={difieren} sin_desglose={sin_desglose}")

# --- AUXILIAR ----------------------------------------------------------
_bloque("auxiliar: el saldo recalculado por el sistema sale marcado en la hoja?")
r = procesar_auxiliar(REAL_PDFS["auxiliar"])
destino = salida / "m6-auxiliar.xlsx"
exportar_auxiliar(r.auxiliar, r.cobertura, destino)
print("hojas y columnas:")
for hoja, cols in _columnas(destino).items():
    print(f"  {hoja:14} {cols}")
origenes: dict[str, int] = {}
for fila in r.auxiliar.filas:
    origenes[fila.saldo_origen] = origenes.get(fila.saldo_origen, 0) + 1
print(f"filas por procedencia del saldo           {origenes}")

# --- BALANZA -----------------------------------------------------------
_bloque("balanza: la fila de Totales que declara el documento sale a la hoja?")
r = procesar_balanza(REAL_PDFS["balanza"])
destino = salida / "m6-balanza.xlsx"
exportar_balanza(r.balanza, r.cobertura, destino)
libro = openpyxl.load_workbook(destino)
hoja = libro["Balanza"]
print("hojas y columnas:")
for h, cols in _columnas(destino).items():
    print(f"  {h:14} {cols}")
print(f"renglones de la hoja (sin encabezado)     {hoja.max_row - 1}")
print(f"filas del parser                          {len(r.balanza.filas)}")
print(f"el parser leyo una fila de Totales        "
      f"{'si' if r.balanza.totales is not None else 'no'}")
if r.balanza.totales is not None:
    nivel1 = [f for f in r.balanza.filas if f.nivel == 1]
    print(f"  declarado  debe {r.balanza.totales.debe} "
          f"haber {r.balanza.totales.haber}")
    print(f"  leido      debe {sum((f.debe for f in nivel1), _CERO)} "
          f"haber {sum((f.haber for f in nivel1), _CERO)} (nivel 1)")
    print("  la fila de Totales NO aparece en la hoja: solo se exportan "
          "`balanza.filas`")
